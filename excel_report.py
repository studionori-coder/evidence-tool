"""撮影したスクリーンショットを、改ページプレビューで整った状態の Excel にまとめる。

単体でも実行できる（セッションフォルダから Excel を作り直す）:
    python excel_report.py <セッションフォルダ>
"""
import datetime
import io
import json
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from PIL import Image as PILImage, ImageOps

# 用紙サイズ (幅mm, 高さmm, Excel の paperSize コード)
PAPERS = {"A4": (210, 297, 9), "A3": (297, 420, 8), "B4": (257, 364, 12)}

MODE_LABELS = {"fullscreen": "全画面", "window": "ウィンドウ", "region": "範囲指定",
               "clipboard": "標準スクショ"}

NCOLS = 8            # エビデンスシートの列数（A〜H）
ROW_PT = 15.0        # 画像エリアの行の高さ(pt)
TITLE_PT = 21.0      # 見出し行の高さ(pt)
INFO_PT = 18.0       # 画面名・メモ行の高さ(pt)
GAP_PT = 9.0         # ブロック間の余白行の高さ(pt)
MARGIN_LR_IN = 0.4   # 左右余白(インチ)
MARGIN_TB_IN = 0.7   # 上下余白(インチ)
HEADER_IN = 0.3      # ヘッダー/フッター位置(インチ)
SAFETY = 0.96        # 1ページに収まるよう高さに持たせる余裕
WIDTH_SAFETY = 0.88  # 画像の横幅に持たせる余裕（列幅の換算誤差対策）

FONT_NAME = "Meiryo UI"
THIN = Side(style="thin", color="808080")
TITLE_FILL = PatternFill("solid", fgColor="DDEBF7")


def _page_geometry(paper, orientation):
    w_mm, h_mm, code = PAPERS.get(paper, PAPERS["A4"])
    if orientation == "landscape":
        w_mm, h_mm = h_mm, w_mm
    printable_w_in = w_mm / 25.4 - MARGIN_LR_IN * 2
    printable_h_in = h_mm / 25.4 - MARGIN_TB_IN * 2
    return code, printable_w_in, printable_h_in


def _place_image(ws, img_path, top_row, area_w_px, area_h_px, col_px, row_px):
    """画像を枠付きで、エリアの中央に縦横比を保って縮小配置する。"""
    with PILImage.open(img_path) as src:
        img = src.convert("RGB")
    img = ImageOps.expand(img, border=2, fill=(128, 128, 128))

    scale = min(area_w_px / img.width, area_h_px / img.height, 1.0)
    disp_w = max(1, int(img.width * scale))
    disp_h = max(1, int(img.height * scale))

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=92)
    buf.seek(0)
    xl_img = XLImage(buf)

    # 横方向は中央寄せ、縦方向は上寄せ（少しだけ余白）
    x = max(0, (int(area_w_px / WIDTH_SAFETY) - disp_w) // 2)
    y = 4
    col, col_off = divmod(x, col_px)
    row, row_off = divmod(y, row_px)
    marker = AnchorMarker(col=int(col), colOff=pixels_to_EMU(col_off),
                          row=top_row - 1 + int(row), rowOff=pixels_to_EMU(row_off))
    xl_img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(
        pixels_to_EMU(disp_w), pixels_to_EMU(disp_h)))
    xl_img.width, xl_img.height = disp_w, disp_h
    ws.add_image(xl_img)


def build_report(entries, base_dir, out_path, title="エビデンス",
                 paper="A4", orientation="portrait", per_page=2):
    """entries: [{no, file, timestamp(ISO), mode, window, comment}, ...]"""
    base_dir = Path(base_dir)
    per_page = max(1, int(per_page))
    paper_code, pw_in, ph_in = _page_geometry(paper, orientation)

    # 印刷可能幅いっぱいに NCOLS 列を並べる（ほぼ等倍で印刷される）
    content_px = int(pw_in * 96 * 0.98)
    col_px = content_px // NCOLS
    content_px = col_px * NCOLS
    col_width = (col_px - 5) / 7.0
    row_px = ROW_PT * 96 / 72

    # 1ページあたり per_page 個のブロック。ブロック = 見出し + 情報 + 画像エリア + 余白
    page_pt = ph_in * 72 * SAFETY
    block_pt = page_pt / per_page
    img_rows = max(5, math.floor((block_pt - TITLE_PT - INFO_PT - GAP_PT) / ROW_PT))
    # 列幅→ピクセルの換算はアプリ（Excel / WPS など）やフォントで数%ずれるため、
    # 横幅は余裕を持たせて枠からはみ出さないようにする
    area_w_px = int(content_px * WIDTH_SAFETY)
    area_h_px = int(img_rows * row_px) - 12

    wb = Workbook()
    ws = wb.active
    ws.title = "エビデンス"
    for c in range(1, NCOLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = col_width

    row = 1
    for i, e in enumerate(entries):
        ts = datetime.datetime.fromisoformat(e["timestamp"])
        mode = MODE_LABELS.get(e.get("mode"), e.get("mode", ""))

        # 見出し行
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
        cell = ws.cell(row=row, column=1,
                       value=f"No.{e['no']}　　撮影日時：{ts:%Y/%m/%d %H:%M:%S}　　（{mode}）")
        cell.font = Font(name=FONT_NAME, size=11, bold=True)
        cell.fill = TITLE_FILL
        cell.alignment = Alignment(vertical="center", indent=1)
        for c in range(1, NCOLS + 1):
            ws.cell(row=row, column=c).border = Border(top=THIN, bottom=THIN,
                                                       left=THIN if c == 1 else None,
                                                       right=THIN if c == NCOLS else None)
        ws.row_dimensions[row].height = TITLE_PT
        title_row = row
        row += 1

        # 画面名・メモ行
        info = f"画面：{e.get('window') or '-'}"
        if e.get("comment"):
            info += f"　／　メモ：{e['comment']}"
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
        cell = ws.cell(row=row, column=1, value=info)
        cell.font = Font(name=FONT_NAME, size=9)
        cell.alignment = Alignment(vertical="center", indent=1, shrink_to_fit=True)
        ws.row_dimensions[row].height = INFO_PT
        row += 1

        # 画像エリア
        for r in range(row, row + img_rows):
            ws.row_dimensions[r].height = ROW_PT
        _place_image(ws, base_dir / e["file"], row, area_w_px, area_h_px, col_px, row_px)
        row += img_rows

        # 余白行
        ws.row_dimensions[row].height = GAP_PT
        e["_cell"] = f"A{title_row}"
        if (i + 1) % per_page == 0 and i + 1 < len(entries):
            ws.row_breaks.append(Break(id=row))
        row += 1

    last_row = max(1, row - 1)

    # 印刷設定：用紙・向き・横1ページに収める・余白・ヘッダー/フッター
    ws.print_area = f"A1:{get_column_letter(NCOLS)}{last_row}"
    ws.page_setup.paperSize = paper_code
    ws.page_setup.orientation = orientation
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=MARGIN_LR_IN, right=MARGIN_LR_IN,
                                  top=MARGIN_TB_IN, bottom=MARGIN_TB_IN,
                                  header=HEADER_IN, footer=HEADER_IN)
    ws.print_options.horizontalCentered = True
    ws.oddHeader.left.text = title
    ws.oddHeader.left.size = 9
    ws.oddHeader.right.text = "出力日：&D"
    ws.oddHeader.right.size = 9
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.size = 9

    # 開いたときに改ページプレビューで表示する
    ws.sheet_view.view = "pageBreakPreview"
    ws.sheet_view.zoomScale = 80
    ws.sheet_view.zoomScaleNormal = 80
    ws.sheet_view.zoomScaleSheetLayoutView = 80
    ws.sheet_view.showGridLines = False

    _build_index(wb, entries, title, paper_code)
    wb.active = 0

    for e in entries:
        e.pop("_cell", None)

    out_path = Path(out_path)
    wb.save(out_path)
    return out_path


def _build_index(wb, entries, title, paper_code):
    """撮影一覧シート（No. からエビデンスへジャンプできる）。"""
    ws = wb.create_sheet("一覧")
    headers = ["No.", "撮影日時", "種別", "画面", "メモ", "ファイル名"]
    widths = [6, 20, 10, 40, 40, 28]
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = TITLE_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws.column_dimensions[get_column_letter(c)].width = w
    for r, e in enumerate(entries, start=2):
        ts = datetime.datetime.fromisoformat(e["timestamp"])
        values = [e["no"], f"{ts:%Y/%m/%d %H:%M:%S}",
                  MODE_LABELS.get(e.get("mode"), e.get("mode", "")),
                  e.get("window", ""), e.get("comment", ""), e["file"]]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            cell.alignment = Alignment(vertical="center", wrap_text=c in (4, 5),
                                       horizontal="center" if c in (1, 3) else None)
        link = ws.cell(row=r, column=1)
        link.hyperlink = f"#'エビデンス'!{e['_cell']}"
        link.font = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")

    last = max(1, len(entries) + 1)
    ws.print_area = f"A1:F{last}"
    ws.print_title_rows = "1:1"
    ws.page_setup.paperSize = paper_code
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.left.text = title
    ws.oddFooter.center.text = "&P / &N"
    ws.freeze_panes = "A2"
    ws.sheet_view.view = "pageBreakPreview"
    ws.sheet_view.zoomScaleSheetLayoutView = 90


def export_session(session_dir, settings=None, out_path=None):
    """セッションフォルダの session.json から Excel を作る。"""
    session_dir = Path(session_dir)
    meta = json.loads((session_dir / "session.json").read_text(encoding="utf-8"))
    entries = meta["entries"]
    settings = settings or {}
    if out_path is None:
        out_path = session_dir / f"エビデンス_{session_dir.name}.xlsx"
    return build_report(entries, session_dir, out_path,
                        title=meta.get("title") or f"エビデンス {session_dir.name}",
                        paper=settings.get("paper", "A4"),
                        orientation=settings.get("orientation", "portrait"),
                        per_page=settings.get("per_page", 2))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("使い方: python excel_report.py <セッションフォルダ>")
        sys.exit(1)
    print(export_session(sys.argv[1]))
